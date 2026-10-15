
import io
import tempfile

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from rest_framework.test import APIClient

from docx import Document
from pypdf import PdfWriter

from resumes.models import Resume


class ResumeUploadTests(TestCase):
    def setUp(self):
        self.temp_media = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_media.cleanup)

        override = override_settings(
            MEDIA_ROOT=self.temp_media.name
        )
        override.enable()
        self.addCleanup(override.disable)

        self.user = User.objects.create_user(
            username="test_resume_user",
            password="TestPassword123!"
        )

        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

        self.url = "/api/resume/upload/"

    def create_pdf(self):
        buffer = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        writer.add_metadata({
            "/Title": "Test CV",
            "/Subject": "Resume test"
        })
        writer.write(buffer)

        # Tambahkan teks menggunakan PDF yang memiliki teks
        # agar fungsi ekstraksi dapat diuji secara realistis.
        from reportlab.pdfgen import canvas

        buffer = io.BytesIO()
        pdf_canvas = canvas.Canvas(buffer)
        pdf_canvas.drawString(50, 750, "Nama: Test User")
        pdf_canvas.drawString(50, 730, "Keahlian: Python dan Django")
        pdf_canvas.save()

        return SimpleUploadedFile(
            "test_cv.pdf",
            buffer.getvalue(),
            content_type="application/pdf"
        )

    def create_docx(self):
        buffer = io.BytesIO()
        document = Document()
        document.add_heading("Curriculum Vitae", level=1)
        document.add_paragraph("Nama: Test User")
        document.add_paragraph("Keahlian: Python dan Django")
        document.save(buffer)

        return SimpleUploadedFile(
            "test_cv.docx",
            buffer.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            )
        )

    def test_upload_pdf_success(self):
        response = self.client.post(
            self.url,
            {"cv_file": self.create_pdf()},
            format="multipart"
        )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data["status"], "success")

        resume = Resume.objects.get(
            id=response.data["data"]["id"]
        )

        self.assertEqual(resume.user, self.user)
        self.assertIn("Test User", resume.extracted_text)
        self.assertTrue(
            resume.cv_file.storage.exists(resume.cv_file.name)
        )

    def test_upload_docx_success(self):
        response = self.client.post(
            self.url,
            {"cv_file": self.create_docx()},
            format="multipart"
        )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data["status"], "success")
        self.assertIn(
            "Test User",
            response.data["data"]["extracted_text"]
        )

    def test_reject_html_file(self):
        upload = SimpleUploadedFile(
            "test.html",
            b"<html><body>Not a CV</body></html>",
            content_type="text/html"
        )

        response = self.client.post(
            self.url,
            {"cv_file": upload},
            format="multipart"
        )

        self.assertEqual(response.status_code, 400)
        self.assertEqual(
            response.data["error"]["code"],
            "INVALID_FILE_TYPE"
        )
        self.assertEqual(Resume.objects.count(), 0)

    def test_reject_file_over_10_mb(self):
        upload = SimpleUploadedFile(
            "large.pdf",
            b"x" * (10 * 1024 * 1024 + 1),
            content_type="application/pdf"
        )

        response = self.client.post(
            self.url,
            {"cv_file": upload},
            format="multipart"
        )

        self.assertEqual(response.status_code, 400)
        self.assertEqual(
            response.data["error"]["code"],
            "FILE_TOO_LARGE"
        )
        self.assertEqual(Resume.objects.count(), 0)

    def test_reject_unauthenticated_upload(self):
        self.client.force_authenticate(user=None)

        response = self.client.post(
            self.url,
            {"cv_file": self.create_docx()},
            format="multipart"
        )

        self.assertEqual(response.status_code, 401)
        self.assertEqual(Resume.objects.count(), 0)

    def test_reject_missing_file(self):
        response = self.client.post(
            self.url,
            {},
            format="multipart"
        )

        self.assertEqual(response.status_code, 400)
        self.assertEqual(
            response.data["error"]["code"],
            "FILE_REQUIRED"
        )
        self.assertEqual(Resume.objects.count(), 0)
